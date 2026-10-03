import openpyxl
import pandas as pd
import json

path = r'c:\Users\PMAL-DATEN\Desktop\projetos\GSVR - prod\util\efetivo - 03-10-2026.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
print("Sheet names:", wb.sheetnames)

sheet = wb.active
print(f"Dimensões: {sheet.max_row} linhas, {sheet.max_column} colunas")

for r in range(1, min(10, sheet.max_row + 1)):
    row_vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
    print(f"Linha {r}: {row_vals}")

df = pd.read_excel(path)
print("\nDataFrame head:")
print(df.head())
print("\nDataFrame columns:")
print(df.columns.tolist())
print(f"\nTotal de linhas no DataFrame: {len(df)}")
