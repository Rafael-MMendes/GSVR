import openpyxl
import json
import paramiko

wb = openpyxl.load_workbook(r"c:\Users\PMAL-DATEN\Desktop\projetos\GSVR - prod\util\novo.xlsx", data_only=True)
ws = wb['GASTOS SETEMBOUTUBRO']

# 1. Pega os dias e meses das colunas D (4) a AH (34)
columns_info = []
for c in range(4, 35):
    col_letter = openpyxl.utils.get_column_letter(c)
    mes = ws.cell(row=12, column=c).value
    dia = ws.cell(row=13, column=c).value
    if dia is not None:
        columns_info.append({
            "col_idx": c,
            "col_letter": col_letter,
            "mes": str(mes) if mes else None,
            "dia": int(dia)
        })

# Preenche os meses para cada coluna
cur_mes = "SETEMBRO"
for item in columns_info:
    if item["dia"] == 1:
        cur_mes = "OUTUBRO"
    item["mes"] = cur_mes
    # gera data ISO para 2026
    m_num = 9 if cur_mes == "SETEMBRO" else 10
    item["data_iso"] = f"2026-{m_num:02d}-{item['dia']:02d}"

# 2. Pega as linhas de turnos (14 a 18)
shifts_info = [
    {"id": "s1", "horario": "14h às 20h", "duracao": "6h", "valor_diaria": 192.03, "row": 14},
    {"id": "s2", "horario": "18h às 00h", "duracao": "6h", "valor_diaria": 192.03, "row": 15},
    {"id": "s3", "horario": "20h às 02h", "duracao": "6h", "valor_diaria": 192.03, "row": 16},
    {"id": "s4", "horario": "14h às 22h", "duracao": "8h", "valor_diaria": 250.00, "row": 17},
    {"id": "s5", "horario": "18h às 02h", "duracao": "8h", "valor_diaria": 250.00, "row": 18},
]

matrix = {}
for s in shifts_info:
    r = s["row"]
    matrix[s["id"]] = {}
    for col in columns_info:
        val = ws.cell(row=r, column=col["col_idx"]).value
        # Se for número, armazena como float/int, se for texto ("E", "L", etc), armazena como string
        if val is None or val == "":
            val_clean = 0
        elif isinstance(val, (int, float)):
            val_clean = int(val) if val == int(val) else val
        else:
            val_clean = str(val).strip()
        matrix[s["id"]][col["data_iso"]] = val_clean

# 3. Totais da planilha
valor_total_ft = float(ws.cell(row=8, column=28).value or 85000.0) # AB8
pm_fora_6h = int(ws.cell(row=20, column=19).value or 0) # S20
pm_fora_8h = int(ws.cell(row=21, column=19).value or 0) # S21

payload = {
    "columns": columns_info,
    "shifts": shifts_info,
    "matrix": matrix,
    "valor_total_ft": valor_total_ft,
    "pm_fora_6h": pm_fora_6h,
    "pm_fora_8h": pm_fora_8h
}

json_str = json.dumps(payload, ensure_ascii=False)
print("Payload gerado com sucesso! Tamanho:", len(json_str))

# 4. Cria a tabela e insere na VPS via SSH
ssh = paramiko.SSHClient()
ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
try:
    ssh.connect('192.168.1.102', username='vps_9bpm', password='asdf1234', timeout=10)
    
    create_sql = """
    CREATE TABLE IF NOT EXISTS PLANILHA_GASTOS_FT (
      id_ciclo INTEGER PRIMARY KEY REFERENCES CICLOS(id_ciclo) ON DELETE CASCADE,
      dados_matriz JSONB NOT NULL,
      valor_total_ft NUMERIC(12,2) DEFAULT 85000.00,
      pm_fora_6h INTEGER DEFAULT 0,
      pm_fora_8h INTEGER DEFAULT 0,
      created_at TIMESTAMP WITH TIME ZONE DEFAULT NOW(),
      updated_at TIMESTAMP WITH TIME ZONE DEFAULT NOW()
    );
    """
    stdin, stdout, stderr = ssh.exec_command(f'docker exec ft-postgres psql -U postgres -d escala_ft -c "{create_sql}"')
    print("Create table:\n" + stdout.read().decode())
    
    # Salva o json temporario no servidor para carregar via psql
    sftp = ssh.open_sftp()
    remote_json_path = "/home/vps_9bpm/gastos_ciclo6.json"
    with sftp.file(remote_json_path, "w") as f:
        f.write(json_str)
    sftp.close()
    
    # Copia para dentro do container e executa upsert
    ssh.exec_command(f"docker cp {remote_json_path} ft-postgres:/tmp/gastos_ciclo6.json")
    
    upsert_sql = """
    INSERT INTO PLANILHA_GASTOS_FT (id_ciclo, dados_matriz, valor_total_ft, pm_fora_6h, pm_fora_8h, updated_at)
    VALUES (6, pg_read_file('/tmp/gastos_ciclo6.json')::jsonb, 85000.00, 0, 0, NOW())
    ON CONFLICT (id_ciclo) DO UPDATE SET
      dados_matriz = EXCLUDED.dados_matriz,
      valor_total_ft = EXCLUDED.valor_total_ft,
      pm_fora_6h = EXCLUDED.pm_fora_6h,
      pm_fora_8h = EXCLUDED.pm_fora_8h,
      updated_at = NOW();
    """
    stdin, stdout, stderr = ssh.exec_command(f'docker exec ft-postgres psql -U postgres -d escala_ft -c "{upsert_sql}"')
    print("Upsert result:\n" + stdout.read().decode())
    
    # Limpa arquivos temporarios
    ssh.exec_command("rm -f /home/vps_9bpm/gastos_ciclo6.json")
    ssh.exec_command("docker exec ft-postgres rm -f /tmp/gastos_ciclo6.json")

finally:
    ssh.close()
