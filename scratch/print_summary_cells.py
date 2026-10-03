import openpyxl

wb = openpyxl.load_workbook(r"c:\Users\PMAL-DATEN\Desktop\projetos\GSVR - prod\util\novo.xlsx", data_only=False)
ws = wb['GASTOS SETEMBOUTUBRO']

for r in range(8, 26):
    for c in range(1, 36):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            col = openpyxl.utils.get_column_letter(c)
            print(f"{col}{r}: {v}")
