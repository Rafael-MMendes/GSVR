import openpyxl

wb = openpyxl.load_workbook(r"c:\Users\PMAL-DATEN\Desktop\projetos\GSVR - prod\util\novo.xlsx", data_only=True)
ws = wb['GASTOS SETEMBOUTUBRO']

for c in range(4, 36):
    col = openpyxl.utils.get_column_letter(c)
    mes = ws.cell(row=12, column=c).value
    dia = ws.cell(row=13, column=c).value
    r14 = ws.cell(row=14, column=c).value
    r15 = ws.cell(row=15, column=c).value
    r16 = ws.cell(row=16, column=c).value
    r17 = ws.cell(row=17, column=c).value
    r18 = ws.cell(row=18, column=c).value
    total_col = sum([x for x in [r14, r15, r16, r17, r18] if isinstance(x, (int, float))])
    print(f"Col {col} ({c}): Mes={mes}, Dia={dia} | 14h-20h: {r14}, 18h-00h: {r15}, 20h-02h: {r16}, 14h-22h: {r17}, 18h-02h: {r18} => Total Dia={total_col}")
