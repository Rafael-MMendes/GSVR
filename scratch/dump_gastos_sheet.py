import openpyxl

wb_formulas = openpyxl.load_workbook(r"c:\Users\PMAL-DATEN\Desktop\projetos\GSVR - prod\util\novo.xlsx", data_only=False)
wb_values = openpyxl.load_workbook(r"c:\Users\PMAL-DATEN\Desktop\projetos\GSVR - prod\util\novo.xlsx", data_only=True)

ws_f = wb_formulas['GASTOS SETEMBOUTUBRO']
ws_v = wb_values['GASTOS SETEMBOUTUBRO']

print(f"Max row: {ws_f.max_row}, Max col: {ws_f.max_column}")

for r in range(1, ws_f.max_row + 1):
    row_data = []
    has_content = False
    for c in range(1, ws_f.max_column + 1):
        cell_f = ws_f.cell(row=r, column=c)
        cell_v = ws_v.cell(row=r, column=c)
        col_letter = openpyxl.utils.get_column_letter(c)
        
        f_val = cell_f.value
        v_val = cell_v.value
        
        if f_val is not None:
            has_content = True
            if str(f_val).startswith('='):
                row_data.append(f"{col_letter}{r}: [F]{f_val} => {v_val}")
            else:
                row_data.append(f"{col_letter}{r}: {f_val}")
    if has_content:
        print(f"--- LINHA {r} ---")
        for item in row_data:
            print("  ", item)
